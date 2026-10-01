from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from src.qc_classifier import SCRIPTING_CATEGORIES, VIDEO_CATEGORIES
from src.version import BUILD_VERSION

VALID_TEAMS = {"scripting", "video_editing", "none"}
VALID_ACTIONS = {"Approve", "Change", "Remove"}


def error_key(row: dict[str, Any]) -> str:
    return f"{row.get('comment_id', '')}:{row.get('issue_index', 0)}:{row.get('team', '')}"


def parse_review_overrides(payload: bytes) -> dict[str, dict[str, Any]]:
    wb = load_workbook(BytesIO(payload), data_only=False)
    if "Review Overrides" not in wb.sheetnames:
        raise ValueError("Reviewed workbook is missing the 'Review Overrides' sheet.")
    if "_Build Info" not in wb.sheetnames:
        raise ValueError("Reviewed workbook is missing build metadata.")
    workbook_build = str(wb["_Build Info"]["B1"].value or "").strip()
    if workbook_build != BUILD_VERSION:
        raise ValueError(
            f"Reviewed workbook build {workbook_build or 'UNKNOWN'} does not match active build {BUILD_VERSION}. "
            "Use the workbook downloaded from the active run."
        )

    ws = wb["Review Overrides"]
    headers = {
        str(cell.value or "").strip(): index
        for index, cell in enumerate(ws[1], start=1)
    }
    required = {
        "Error Key",
        "Review Action",
        "Manual Team",
        "Manual Category",
        "Manual Assignee",
        "Manual Count",
        "Reviewer Notes",
    }
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError(
            "Reviewed workbook has an incompatible Review Overrides layout: "
            + ", ".join(missing)
        )

    overrides: dict[str, dict[str, Any]] = {}
    for row_number in range(2, ws.max_row + 1):
        key = str(ws.cell(row_number, headers["Error Key"]).value or "").strip()
        action = str(ws.cell(row_number, headers["Review Action"]).value or "").strip()
        if not key or not action:
            continue
        if action not in VALID_ACTIONS:
            raise ValueError(
                f"Invalid Review Action '{action}' on row {row_number}. "
                "Use Approve, Change, or Remove."
            )

        manual_count_raw = ws.cell(row_number, headers["Manual Count"]).value
        manual_count: int | None = None
        if manual_count_raw not in (None, ""):
            try:
                manual_count = int(manual_count_raw)
            except (TypeError, ValueError) as exc:
                raise ValueError(
                    f"Manual Count must be a whole number on row {row_number}."
                ) from exc
            if manual_count < 0:
                raise ValueError(
                    f"Manual Count cannot be negative on row {row_number}."
                )

        override = {
            "action": action,
            "manual_team": str(ws.cell(row_number, headers["Manual Team"]).value or "").strip(),
            "manual_category": str(ws.cell(row_number, headers["Manual Category"]).value or "").strip(),
            "manual_assignee": str(ws.cell(row_number, headers["Manual Assignee"]).value or "").strip(),
            "manual_count": manual_count,
            "reviewer_notes": str(ws.cell(row_number, headers["Reviewer Notes"]).value or "").strip(),
        }

        if key in overrides and overrides[key] != override:
            raise ValueError(f"Conflicting duplicate review action for Error Key {key}.")
        overrides[key] = override

    return overrides


def merge_review_overrides(*maps: dict[str, dict[str, Any]]) -> dict[str, dict[str, Any]]:
    merged: dict[str, dict[str, Any]] = {}
    for mapping in maps:
        for key, override in mapping.items():
            if key in merged and merged[key] != override:
                raise ValueError(
                    f"Conflicting review decisions were supplied for Error Key {key}."
                )
            merged[key] = override
    return merged


def _bundle_verified_map(bundles: list[dict[str, Any]]) -> dict[str, bool]:
    result: dict[str, bool] = {}
    for bundle in bundles:
        project = bundle.get("project") or {}
        article_id = str(project.get("article_id") or "")
        resolution = bundle.get("resolution")
        audit = bundle.get("extraction_audit") or {}
        result[article_id] = (
            resolution is not None
            and getattr(resolution, "status", "") == "resolved"
            and audit.get("status") == "verified"
        )
    return result


def _category_valid(team: str, category: str) -> bool:
    if team == "scripting":
        return category in SCRIPTING_CATEGORIES
    if team == "video_editing":
        return category in VIDEO_CATEGORIES
    return category in {"", "None"}


def apply_review_overrides(
    error_rows: list[dict[str, Any]],
    bundles: list[dict[str, Any]],
    overrides: dict[str, dict[str, Any]],
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    rows = deepcopy(error_rows)
    verified = _bundle_verified_map(bundles)
    applied = 0
    approved = 0
    changed = 0
    removed = 0
    known_keys = {error_key(row) for row in rows}

    unknown = sorted(set(overrides) - known_keys)
    if unknown:
        raise ValueError(
            f"{len(unknown)} reviewed Error Key(s) do not belong to the active QC run."
        )

    for row in rows:
        key = error_key(row)
        override = overrides.get(key)
        if not override:
            continue

        action = override["action"]
        original_team = str(row.get("team") or "")
        original_category = str(row.get("ai_category") or "")
        original_assignee = str(row.get("ai_assignee") or "")
        original_count = int(row.get("ai_error_count") or 0)

        row["original_team"] = original_team
        row["original_ai_category"] = original_category
        row["original_ai_assignee"] = original_assignee
        row["original_ai_error_count"] = original_count
        row["review_action"] = action
        row["reviewer_notes"] = override.get("reviewer_notes", "")
        row["review_applied"] = True

        if action == "Approve":
            if original_team == "review":
                raise ValueError(
                    f"Error Key {key} came from a failed/unknown classification. "
                    "Use Change or Remove instead of Approve."
                )
            approved += 1

        elif action == "Change":
            team = override.get("manual_team") or original_team
            category = override.get("manual_category") or original_category
            assignee = override.get("manual_assignee") or original_assignee
            count = (
                override["manual_count"]
                if override.get("manual_count") is not None
                else original_count
            )

            if team not in VALID_TEAMS:
                raise ValueError(f"Invalid Manual Team '{team}' for Error Key {key}.")
            if not _category_valid(team, category):
                raise ValueError(
                    f"Category '{category}' is not valid for team '{team}' on Error Key {key}."
                )
            if team in {"scripting", "video_editing"} and count <= 0:
                raise ValueError(
                    f"Manual Count must be at least 1 for retained Error Key {key}."
                )

            row["team"] = team
            row["ai_category"] = category
            row["ai_assignee"] = assignee
            row["ai_error_count"] = count
            changed += 1

        elif action == "Remove":
            row["team"] = "none"
            row["ai_category"] = "None"
            row["ai_assignee"] = ""
            row["ai_error_count"] = 0
            removed += 1

        article_id = str(row.get("article_id") or "")
        current_team = str(row.get("team") or "")
        current_assignee = str(row.get("ai_assignee") or "")

        if current_team in {"scripting", "video_editing"}:
            row["needs_review"] = False
            row["performance_eligible"] = bool(
                verified.get(article_id, False) and current_assignee
            )
            if not verified.get(article_id, False):
                row["needs_review"] = True
                row["performance_eligible"] = False
                row["classification_reason"] = (
                    str(row.get("classification_reason") or "")
                    + " | Reviewer decision recorded, but extraction remains unverified."
                ).strip(" |")
        else:
            row["needs_review"] = False
            row["performance_eligible"] = False

        row["manual_team"] = override.get("manual_team", "")
        row["manual_category"] = override.get("manual_category", "")
        row["manual_assignee"] = override.get("manual_assignee", "")
        row["manual_count"] = override.get("manual_count")
        applied += 1

    agreement_denominator = approved + changed + removed
    agreement = (
        approved / agreement_denominator
        if agreement_denominator
        else None
    )

    return rows, {
        "applied": applied,
        "approved": approved,
        "changed": changed,
        "removed": removed,
        "reviewed_ai_agreement": agreement,
    }
