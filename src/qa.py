from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from src.qc_classifier import OPERATIONAL_CATEGORIES, SCRIPTING_CATEGORIES, VIDEO_CATEGORIES
from src.report_generator import _performance_eligible_article_ids
from src.version import BUILD_VERSION


def _formula_count(wb) -> int:
    count = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    count += 1
    return count


def _sheet_headers(ws) -> list[str]:
    return [str(cell.value or "") for cell in ws[1]]


def _expected_main_headers(team: str) -> list[str]:
    categories = (
        SCRIPTING_CATEGORIES
        if team == "scripting"
        else {**VIDEO_CATEGORIES, "VO Redos": OPERATIONAL_CATEGORIES["VO Redos"]}
    )
    contributor = "Writer" if team == "scripting" else "Video Editor"
    date_header = "Date Script Draft Sent" if team == "scripting" else "Date Video Released"
    headers = ["Article Number", contributor, date_header]
    for category in categories:
        headers.extend([category, "QC Comment"])
    headers.append("Total Errors" if team == "scripting" else "Total Editor Errors")
    if team == "video_editing":
        headers.append("Total QC Items")
    return headers

def _main_is_sorted(ws) -> bool:
    values = []
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row, 3).value
        values.append((value, str(ws.cell(row, 1).value or "")))
    return values == sorted(values, key=lambda item: (item[0], item[1]))


def _review_headers_ok(ws) -> bool:
    required = {
        "Error Key",
        "Article Number",
        "Version",
        "Timecode",
        "Frame.io Review Link",
        "Comment",
        "AI Team",
        "AI Category",
        "AI Assignee",
        "AI Count",
        "AI Needs Review",
        "AI Performance Eligible",
        "Review Action",
        "Manual Team",
        "Manual Category",
        "Manual Assignee",
        "Manual Count",
        "Reviewer Notes",
    }
    return required.issubset(set(_sheet_headers(ws)))


def audit_workbook(
    payload: bytes,
    *,
    team: str,
    projects: list[dict[str, Any]],
    bundles: list[dict[str, Any]],
    error_rows: list[dict[str, Any]],
) -> dict[str, Any]:
    wb = load_workbook(BytesIO(payload), data_only=False)
    categories = (
        SCRIPTING_CATEGORIES
        if team == "scripting"
        else {**VIDEO_CATEGORIES, "VO Redos": OPERATIONAL_CATEGORIES["VO Redos"]}
    )
    main_sheet = "QC by Draft Date" if team == "scripting" else "Video QC by Release Date"

    required = {
        main_sheet,
        "Description",
        "Monthly Trend",
        "Analysis Guide",
        "Extraction Audit",
        "Needs Review",
        "Review Overrides",
        "Comment Coverage",
        "Non-Performance Items",
        "Excluded Version Policy",
        "Error Detail",
        "Raw Frame.io Comments",
        "_Build Info",
    }
    required.add("Writer Summary" if team == "scripting" else "Editor Summary")
    required.add("Writer Monthly Trend" if team == "scripting" else "Editor Monthly Trend")
    if team == "video_editing":
        required.update({"Common Error Patterns", "Frame.io Version Breakdown", "Scope Reconciliation"})

    checks: dict[str, bool] = {}
    notes: list[str] = []

    checks["required_sheets"] = required.issubset(set(wb.sheetnames))
    if not checks["required_sheets"]:
        notes.append("One or more required workbook tabs are missing.")

    if main_sheet in wb.sheetnames:
        ws = wb[main_sheet]
        checks["main_headers"] = _sheet_headers(ws) == _expected_main_headers(team)
        checks["project_row_count"] = ws.max_row - 1 == len(projects)
        checks["main_date_sort"] = _main_is_sorted(ws)

        if team == "scripting":
            expected_total = sum(
                int(row.get("ai_error_count") or 0)
                for row in error_rows
                if row.get("team") == "scripting"
            )
            actual_total = sum(
                int(ws.cell(row, ws.max_column).value or 0)
                for row in range(2, ws.max_row + 1)
            )
            checks["main_total_consistency"] = actual_total == expected_total
        else:
            expected_editor = sum(
                int(row.get("ai_error_count") or 0)
                for row in error_rows
                if row.get("team") == "video_editing"
            )
            expected_vo = sum(
                int(row.get("ai_error_count") or 0)
                for row in error_rows
                if row.get("team") == "operational" and row.get("ai_category") == "VO Redos"
            )
            actual_editor = sum(
                int(ws.cell(row, ws.max_column - 1).value or 0)
                for row in range(2, ws.max_row + 1)
            )
            actual_qc = sum(
                int(ws.cell(row, ws.max_column).value or 0)
                for row in range(2, ws.max_row + 1)
            )
            checks["main_total_consistency"] = (
                actual_editor == expected_editor
                and actual_qc == expected_editor + expected_vo
            )

        if not checks["main_total_consistency"]:
            notes.append(f"{team} main-sheet totals do not reconcile to classified rows.")

        by_article = {
            str(ws.cell(row_number, 1).value or ""): row_number
            for row_number in range(2, ws.max_row + 1)
        }
        verbatim_ok = True
        category_order = list(categories)
        for row in error_rows:
            if team == "scripting":
                belongs = row.get("team") == "scripting"
            else:
                belongs = (
                    row.get("team") == "video_editing"
                    or (row.get("team") == "operational" and row.get("ai_category") == "VO Redos")
                )
            if not belongs:
                continue
            article = str(row.get("item_name") or row.get("article_id") or "")
            category = str(row.get("ai_category") or "")
            source = str(row.get("comment_text") or "")
            if not source or article not in by_article or category not in category_order:
                continue
            excel_row = by_article[article]
            comment_col = 5 + (2 * category_order.index(category))
            displayed = str(ws.cell(excel_row, comment_col).value or "")
            if source not in displayed:
                verbatim_ok = False
                break
        checks["main_verbatim_source_comments"] = verbatim_ok
        if not verbatim_ok:
            notes.append("At least one main-sheet QC comment does not contain the exact Frame.io source comment.")
    else:
        checks["main_headers"] = False
        checks["project_row_count"] = False
        checks["main_date_sort"] = False
        checks["main_total_consistency"] = False
        checks["main_verbatim_source_comments"] = False

    if "Description" in wb.sheetnames:
        desc = wb["Description"]
        observed = {
            str(desc.cell(row, 1).value or ""): str(desc.cell(row, 2).value or "")
            for row in range(2, desc.max_row + 1)
            if desc.cell(row, 1).value
        }
        checks["category_definitions"] = all(
            observed.get(category) == definition
            for category, definition in categories.items()
        )
    else:
        checks["category_definitions"] = False

    checks["no_formulas"] = _formula_count(wb) == 0

    if "_Build Info" in wb.sheetnames:
        checks["build_metadata"] = wb["_Build Info"]["B1"].value == BUILD_VERSION
    else:
        checks["build_metadata"] = False

    if "Extraction Audit" in wb.sheetnames:
        audit_ws = wb["Extraction Audit"]
        checks["extraction_audit_row_count"] = audit_ws.max_row - 1 == len(projects)
        link_ok = all(
            not str(audit_ws.cell(row, 2).value or "")
            or bool(audit_ws.cell(row, 2).hyperlink)
            for row in range(2, audit_ws.max_row + 1)
        )
        checks["extraction_links"] = link_ok
    else:
        checks["extraction_audit_row_count"] = False
        checks["extraction_links"] = False

    included_comments = [
        comment
        for bundle in bundles
        for comment in (bundle.get("comments") or [])
    ]
    excluded_comments = []

    if "Comment Coverage" in wb.sheetnames:
        coverage = wb["Comment Coverage"]
        checks["comment_coverage_row_count"] = coverage.max_row - 1 == len(included_comments)
        source_values = [
            str(coverage.cell(row, 6).value or "")
            for row in range(2, coverage.max_row + 1)
        ]
        checks["comment_coverage_verbatim"] = sorted(source_values) == sorted(
            str(comment.get("text") or "") for comment in included_comments
        )
    else:
        checks["comment_coverage_row_count"] = False
        checks["comment_coverage_verbatim"] = False

    if "Excluded Version Policy" in wb.sheetnames:
        excluded = wb["Excluded Version Policy"]
        expected_excluded_versions = sum(
            min(2, len(bundle.get("versions") or []))
            for bundle in bundles
            if bundle.get("versions")
        )
        checks["excluded_version_policy_rows"] = excluded.max_row - 1 == expected_excluded_versions
    else:
        checks["excluded_version_policy_rows"] = False

    if "Needs Review" in wb.sheetnames:
        checks["needs_review_structure"] = wb["Needs Review"].max_column >= 11
    else:
        checks["needs_review_structure"] = False

    if "Review Overrides" in wb.sheetnames:
        checks["review_workflow_structure"] = _review_headers_ok(wb["Review Overrides"])
    else:
        checks["review_workflow_structure"] = False

    if "Monthly Trend" in wb.sheetnames:
        trend_header = str(wb["Monthly Trend"]["A1"].value or "")
        checks["trend_scope_labelled"] = "selected" in trend_header.casefold()
    else:
        checks["trend_scope_labelled"] = False

    if team == "video_editing" and "Common Error Patterns" in wb.sheetnames:
        patterns = wb["Common Error Patterns"]
        observed_total = sum(
            int(patterns.cell(row, 3).value or 0)
            for row in range(2, patterns.max_row + 1)
        )
        expected_total = sum(
            int(row.get("ai_error_count") or 0)
            for row in error_rows
            if row.get("team") == "video_editing"
        )
        checks["pattern_total_reconciles"] = observed_total == expected_total
        checks["pattern_recurring_flag"] = all(
            str(patterns.cell(row, 6).value or "") in {"Yes", "No"}
            for row in range(2, patterns.max_row + 1)
        )
    else:
        checks["pattern_total_reconciles"] = True
        checks["pattern_recurring_flag"] = True

    eligible_ids = _performance_eligible_article_ids(projects, bundles, error_rows, team)
    summary_name = "Writer Summary" if team == "scripting" else "Editor Summary"
    if summary_name in wb.sheetnames:
        summary = wb[summary_name]
        workbook_articles = sum(
            int(summary.cell(row, 2).value or 0)
            for row in range(2, summary.max_row + 1)
        )
        checks["performance_denominator_consistency"] = workbook_articles == len(eligible_ids)
    else:
        checks["performance_denominator_consistency"] = False

    passed = sum(1 for value in checks.values() if value)
    score = passed / max(len(checks), 1)
    return {
        "team": team,
        "score": score,
        "checks": checks,
        "notes": notes,
        "performance_eligible_articles": len(eligible_ids),
    }

def audit_run(
    *,
    projects: list[dict[str, Any]],
    bundles: list[dict[str, Any]],
    classifications: list[dict[str, Any]],
    error_rows: list[dict[str, Any]],
    scripting_report: bytes,
    video_report: bytes,
) -> dict[str, Any]:
    comments = [
        comment
        for bundle in bundles
        for comment in (bundle.get("comments") or [])
    ]
    excluded_comments = [
        comment
        for bundle in bundles
        for comment in (bundle.get("excluded_comments") or [])
    ]
    source_comments = comments + excluded_comments

    resolved_projects = [
        bundle
        for bundle in bundles
        if bundle.get("resolution") is not None
        and getattr(bundle["resolution"], "status", "") == "resolved"
    ]
    extraction_verified = [
        bundle
        for bundle in bundles
        if (bundle.get("extraction_audit") or {}).get("status") == "verified"
    ]

    classification_failures = [
        row for row in classifications if row.get("_classification_error")
    ]

    commenter_count = sum(
        1 for row in comments if str(row.get("commenter") or "").strip()
    )

    comment_ids = [
        str(row.get("comment_id") or "")
        for row in source_comments
        if row.get("comment_id")
    ]

    zero_error_policy_ok = True
    version_exclusion_ok = True
    for bundle in bundles:
        audit = bundle.get("extraction_audit") or {}
        bundle_comments = bundle.get("comments") or []
        if not bundle_comments and audit.get("zero_error_verified"):
            if audit.get("status") != "verified" or int(audit.get("eligible_versions") or 0) <= 0:
                zero_error_policy_ok = False

        versions = bundle.get("versions") or []
        for index, version in enumerate(versions):
            expected = 0 < index < len(versions) - 1
            if bool(version.get("included_for_qc")) != expected:
                version_exclusion_ok = False
                break

        included_versions = {
            int(comment.get("version_number") or 0)
            for comment in bundle.get("comments") or []
        }
        if 1 in included_versions:
            version_exclusion_ok = False
        if versions and len(versions) in included_versions:
            version_exclusion_ok = False

    unverified_article_ids = {
        str((bundle.get("project") or {}).get("article_id") or "")
        for bundle in bundles
        if (bundle.get("extraction_audit") or {}).get("status") != "verified"
    }
    scripting_eligible = _performance_eligible_article_ids(
        projects, bundles, error_rows, "scripting"
    )
    video_eligible = _performance_eligible_article_ids(
        projects, bundles, error_rows, "video_editing"
    )
    unverified_excluded = not (
        unverified_article_ids & scripting_eligible
        or unverified_article_ids & video_eligible
    )

    invalid_team_category_rows = [
        row
        for row in error_rows
        if (
            row.get("team") == "scripting"
            and row.get("ai_category") not in SCRIPTING_CATEGORIES
        )
        or (
            row.get("team") == "video_editing"
            and row.get("ai_category") not in VIDEO_CATEGORIES
        )
        or (
            row.get("team") == "operational"
            and row.get("ai_category") not in OPERATIONAL_CATEGORIES
        )
    ]

    data_checks = {
        "projects_present": bool(projects),
        "frameio_resolution_complete": len(resolved_projects) == len(projects),
        "extraction_verification_complete": len(extraction_verified) == len(projects),
        "classification_coverage_complete": len(classification_failures) == 0,
        "comment_ids_unique": len(comment_ids) == len(set(comment_ids)),
        "version_exclusion_policy": version_exclusion_ok,
        "zero_error_policy": zero_error_policy_ok,
        "unverified_excluded_from_performance": unverified_excluded,
        "team_category_consistency": len(invalid_team_category_rows) == 0,
        "commenter_identity_coverage": (
            (commenter_count / len(comments)) if comments else 1.0
        ),
    }

    scripting = audit_workbook(
        scripting_report,
        team="scripting",
        projects=projects,
        bundles=bundles,
        error_rows=error_rows,
    )
    video = audit_workbook(
        video_report,
        team="video_editing",
        projects=projects,
        bundles=bundles,
        error_rows=error_rows,
    )

    binary_checks = [
        data_checks["projects_present"],
        data_checks["frameio_resolution_complete"],
        data_checks["extraction_verification_complete"],
        data_checks["classification_coverage_complete"],
        data_checks["comment_ids_unique"],
        data_checks["version_exclusion_policy"],
        data_checks["zero_error_policy"],
        data_checks["unverified_excluded_from_performance"],
        data_checks["team_category_consistency"],
        *scripting["checks"].values(),
        *video["checks"].values(),
    ]
    structural_score = sum(1 for value in binary_checks if value) / max(len(binary_checks), 1)

    critical_pass = (
        data_checks["frameio_resolution_complete"]
        and data_checks["classification_coverage_complete"]
        and data_checks["version_exclusion_policy"]
        and data_checks["zero_error_policy"]
        and data_checks["unverified_excluded_from_performance"]
        and data_checks["team_category_consistency"]
    )

    notes = [*scripting["notes"], *video["notes"]]
    if invalid_team_category_rows:
        examples = ", ".join(
            f"{row.get('item_name') or row.get('article_id')}: {row.get('team')} / {row.get('ai_category')}"
            for row in invalid_team_category_rows[:5]
        )
        notes.append(
            f"{len(invalid_team_category_rows)} invalid team/category pair(s) were detected: {examples}."
        )

    missing_scriptwriters = [
        project.get("item_name") or project.get("article_id")
        for project in projects
        if not str(project.get("scriptwriter") or "").strip()
    ]
    video_handovers = [
        project.get("item_name") or project.get("article_id")
        for project in projects
        if len({
            str(project.get("science_video_editor") or "").strip(),
            str(project.get("finishing_editor") or "").strip(),
            str(project.get("rough_editor") or "").strip(),
        } - {""}) > 1
    ]
    if missing_scriptwriters:
        notes.append(
            f"{len(missing_scriptwriters)} project(s) have no Scriptwriter in Monday.com and are excluded from individual writer performance: "
            + ", ".join(missing_scriptwriters[:10])
            + ("." if len(missing_scriptwriters) <= 10 else ", ...")
        )
    if video_handovers:
        notes.append(
            f"{len(video_handovers)} project(s) have multiple video editors in Monday.com and require individual attribution review: "
            + ", ".join(video_handovers[:10])
            + ("." if len(video_handovers) <= 10 else ", ...")
        )
    if classification_failures:
        notes.append(
            f"{len(classification_failures)} comment(s) remain unclassified after recovery."
        )
    if len(resolved_projects) != len(projects):
        notes.append(
            f"{len(projects) - len(resolved_projects)} project(s) remain unresolved in Frame.io."
        )
    if len(extraction_verified) != len(projects):
        notes.append(
            f"{len(projects) - len(extraction_verified)} project(s) do not have fully verified intermediate-version extraction and are excluded from performance denominators."
        )

    commenter_score = float(data_checks["commenter_identity_coverage"])
    if comments and commenter_score < 0.95:
        notes.append(
            f"Commenter identity coverage is {commenter_score:.1%}; this is audit metadata only and does not count as semantic classification accuracy."
        )

    notes.append(
        "Structural QA does not measure semantic classification accuracy. Semantic accuracy is established through the Review Overrides workflow and reviewer agreement."
    )

    return {
        "score": structural_score,
        "score_type": "structural_and_coverage",
        "target": 0.95,
        "passed": structural_score >= 0.95 and critical_pass,
        "critical_pass": critical_pass,
        "semantic_accuracy": None,
        "data_checks": data_checks,
        "scripting": scripting,
        "video": video,
        "notes": notes,
    }
