from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from src.qc_classifier import OPERATIONAL_CATEGORIES, SCRIPTING_CATEGORIES, VIDEO_CATEGORIES
from src.version import BUILD_LABEL, BUILD_VERSION

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
TOTAL_FILL = PatternFill("solid", fgColor="D9EAF7")
TOTAL_FONT = Font(color="17365D", bold=True)
CAUTION_FILL = PatternFill("solid", fgColor="FCE4D6")
ERROR_FILL = PatternFill("solid", fgColor="F4CCCC")
INPUT_FONT = Font(color="0000FF")
IMPORTED_FONT = Font(color="008000")
STATIC_FONT = Font(color="666666")
THIN = Side(style="thin", color="D9E2F3")
MEDIUM = Side(style="medium", color="D9E2F3")
BOTTOM_BORDER = Border(bottom=THIN)
MAIN_HEADER_BORDER = Border(bottom=MEDIUM)


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%b %d, %Y", "%b %d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def _month_key(value: Any) -> str:
    parsed = _as_date(value)
    return parsed.strftime("%Y-%m") if parsed else "Unknown"


def _style_header(ws, headers: list[str]) -> None:
    for idx, header in enumerate(headers, 1):
        cell = ws.cell(1, idx, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BOTTOM_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 42
    ws.sheet_view.showGridLines = False


def _style_main_header(ws, headers: list[str]) -> None:
    for idx, header in enumerate(headers, 1):
        cell = ws.cell(1, idx, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = MAIN_HEADER_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 42
    ws.sheet_view.showGridLines = False


def _make_link(cell, url: str) -> None:
    if not url:
        return
    cell.hyperlink = url
    cell.style = "Hyperlink"


def _unique_texts(rows: list[dict[str, Any]]) -> list[str]:
    result: list[str] = []
    for row in rows:
        # QC Comment must always be the exact Frame.io source comment.
        # Never use AI-generated/paraphrased wording in a user-facing comment field.
        text = str(row.get("comment_text") or "")
        if text.strip() and text not in result:
            result.append(text)
    return result


def _distinct_editors(project: dict[str, Any]) -> list[str]:
    values = [
        str(project.get("science_video_editor") or "").strip(),
        str(project.get("finishing_editor") or "").strip(),
        str(project.get("rough_editor") or "").strip(),
    ]
    return [value for index, value in enumerate(values) if value and value not in values[:index]]


def _article_contributor(project: dict[str, Any], team: str) -> str:
    if team == "scripting":
        return str(project.get("scriptwriter") or "").strip()
    editors = _distinct_editors(project)
    if len(editors) == 1:
        return editors[0]
    if not editors:
        return ""
    return "Needs Review"


def _category_config(team: str) -> dict[str, str]:
    if team == "scripting":
        return SCRIPTING_CATEGORIES
    return {**VIDEO_CATEGORIES, "VO Redos": OPERATIONAL_CATEGORIES["VO Redos"]}


def _performance_category_config(team: str) -> dict[str, str]:
    return SCRIPTING_CATEGORIES if team == "scripting" else VIDEO_CATEGORIES


def _row_in_report(row: dict[str, Any], team: str) -> bool:
    if team == "scripting":
        return row.get("team") == "scripting"
    return (
        row.get("team") == "video_editing"
        or (row.get("team") == "operational" and row.get("ai_category") == "VO Redos")
    )


def _team_config(team: str) -> dict[str, str]:
    if team == "scripting":
        return {
            "main_sheet": "QC by Draft Date",
            "contributor": "Writer",
            "date_header": "Date Script Draft Sent",
            "date_key": "draft_script_sent",
            "summary": "Writer Summary",
            "contributor_trend": "Writer Monthly Trend",
            "trend_month_header": "Draft Month (selected release cohort)",
        }
    return {
        "main_sheet": "Video QC by Release Date",
        "contributor": "Video Editor",
        "date_header": "Date Video Released",
        "date_key": "date_video_released",
        "summary": "Editor Summary",
        "contributor_trend": "Editor Monthly Trend",
        "trend_month_header": "Release Month (selected cohort)",
    }


def _bundle_map(bundles: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {
        str((bundle.get("project") or {}).get("article_id") or ""): bundle
        for bundle in bundles
    }


def _extraction_verified(bundle: dict[str, Any] | None) -> bool:
    if not bundle:
        return False
    resolution = bundle.get("resolution")
    audit = bundle.get("extraction_audit") or {}
    return (
        resolution is not None
        and getattr(resolution, "status", "") == "resolved"
        and audit.get("status") == "verified"
    )


def _performance_eligible_article_ids(
    projects: list[dict[str, Any]],
    bundles: list[dict[str, Any]],
    error_rows: list[dict[str, Any]],
    team: str,
) -> set[str]:
    bundle_by_article = _bundle_map(bundles)
    blocked_by_article: set[str] = set()

    for row in error_rows:
        article_id = str(row.get("article_id") or "")
        row_team = row.get("team")
        if row_team == "review":
            blocked_by_article.add(article_id)
        elif row_team == team and not row.get("performance_eligible", True):
            blocked_by_article.add(article_id)

    eligible: set[str] = set()
    for project in projects:
        article_id = str(project.get("article_id") or "")
        contributor = _article_contributor(project, team)
        if not contributor or contributor == "Needs Review":
            continue
        if not _extraction_verified(bundle_by_article.get(article_id)):
            continue
        if article_id in blocked_by_article:
            continue
        eligible.add(article_id)
    return eligible


def _team_trend_eligible_article_ids(
    projects: list[dict[str, Any]],
    bundles: list[dict[str, Any]],
    error_rows: list[dict[str, Any]],
) -> set[str]:
    bundle_by_article = _bundle_map(bundles)
    blocked_by_article = {
        str(row.get("article_id") or "")
        for row in error_rows
        if row.get("team") == "review"
    }
    return {
        str(project.get("article_id") or "")
        for project in projects
        if _extraction_verified(
            bundle_by_article.get(str(project.get("article_id") or ""))
        )
        and str(project.get("article_id") or "") not in blocked_by_article
    }


def _main_sheet(wb, projects, bundles, error_rows, team):
    cfg = _team_config(team)
    categories = _category_config(team)
    ws = wb.create_sheet(cfg["main_sheet"])
    headers = ["Article Number", cfg["contributor"], cfg["date_header"]]
    for category in categories:
        headers.extend([category, "QC Comment"])
    headers.append("Total Errors" if team == "scripting" else "Total Editor Errors")
    if team == "video_editing":
        headers.append("Total QC Items")
    _style_main_header(ws, headers)

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 23
    ws.column_dimensions["C"].width = 20
    category_widths = [13, 23, 30, 20]
    comment_widths = [53, 48, 48, 48]
    col = 4
    for index, _ in enumerate(categories):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = category_widths[min(index, len(category_widths) - 1)]
        ws.column_dimensions[ws.cell(1, col + 1).column_letter].width = comment_widths[min(index, len(comment_widths) - 1)]
        col += 2
    ws.column_dimensions[ws.cell(1, col).column_letter].width = 14

    by_article: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in error_rows:
        if _row_in_report(row, team):
            by_article[str(row.get("article_id") or "")].append(row)

    bundle_by_article = _bundle_map(bundles)
    sorted_projects = sorted(
        projects,
        key=lambda project: (
            _as_date(project.get(cfg["date_key"])) or date.max,
            str(project.get("article_id") or ""),
            str(project.get("item_name") or ""),
        ),
    )

    for project in sorted_projects:
        article_id = str(project.get("article_id") or "")
        rows = by_article.get(article_id, [])
        values: list[Any] = [
            project.get("item_name") or article_id,
            _article_contributor(project, team),
            _as_date(project.get(cfg["date_key"])),
        ]
        total_editor_errors = 0
        total_qc_items = 0
        for category in categories:
            category_rows = [row for row in rows if row.get("ai_category") == category]
            count = sum(int(row.get("ai_error_count") or 0) for row in category_rows)
            total_qc_items += count
            if team == "scripting" or category in VIDEO_CATEGORIES:
                total_editor_errors += count
            comments = _unique_texts(category_rows)
            values.extend([count, "\n\n".join(comments) if comments else None])
        values.append(total_editor_errors)
        if team == "video_editing":
            values.append(total_qc_items)
        ws.append(values)

        row_number = ws.max_row
        ws.row_dimensions[row_number].height = 72
        ws.cell(row_number, 3).number_format = "yyyy-mm-dd"
        for column in range(1, len(headers) + 1):
            cell = ws.cell(row_number, column)
            cell.border = BOTTOM_BORDER
            if column in (1, 2, 3):
                cell.alignment = Alignment(vertical="top")
        for column in range(4, len(headers), 2):
            ws.cell(row_number, column).alignment = Alignment(horizontal="center", vertical="top")
            if column + 1 < len(headers):
                ws.cell(row_number, column + 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row_number, len(headers)).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row_number, len(headers)).fill = TOTAL_FILL
        ws.cell(row_number, len(headers)).font = TOTAL_FONT

        bundle = bundle_by_article.get(article_id)
        if not _extraction_verified(bundle):
            audit_note = str((bundle or {}).get("extraction_audit", {}).get("note") or "Extraction not verified.")
            ws.cell(row_number, 1).fill = CAUTION_FILL
            ws.cell(row_number, len(headers)).fill = CAUTION_FILL
            ws.cell(row_number, 1).comment = Comment(
                "This article is shown for completeness but is excluded from performance denominators until extraction is verified. "
                + audit_note,
                "QC Automation",
            )

        contributor = _article_contributor(project, team)
        if not contributor:
            ws.cell(row_number, 2).fill = CAUTION_FILL
            ws.cell(row_number, 2).comment = Comment(
                "Contributor is missing in Monday.com. This article is excluded from individual contributor performance until the source assignment is corrected.",
                "QC Automation",
            )
        elif team == "video_editing" and contributor == "Needs Review":
            ws.cell(row_number, 2).fill = CAUTION_FILL
            ws.cell(row_number, 2).comment = Comment(
                "Multiple video editors are assigned in Monday.com. Article-level QC remains visible, but individual editor attribution requires review.",
                "QC Automation",
            )

    return ws


def _description_sheet(wb, team):
    categories = _category_config(team)
    ws = wb.create_sheet("Description")
    ws.append(["Error Category", "Description"])
    if team == "scripting":
        order = [
            "Scientific Accuracy & Completeness",
            "Language & Narration",
            "Titles & On-screen Text",
            "Scientific Formatting",
        ]
    else:
        order = list(categories)
    for category in order:
        ws.append([category, categories[category]])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 95
    return ws


def _article_error_map(error_rows, team, performance_only=False):
    by_article = defaultdict(list)
    for row in error_rows:
        if performance_only:
            if row.get("team") != team or not row.get("performance_eligible", False):
                continue
        elif not _row_in_report(row, team):
            continue
        by_article[str(row.get("article_id") or "")].append(row)
    return by_article


def _monthly_trend_sheet(wb, projects, bundles, error_rows, team):
    cfg = _team_config(team)
    categories = _category_config(team)
    ws = wb.create_sheet("Monthly Trend")
    headers = [cfg["trend_month_header"], "Verified Articles", *categories.keys(), "Total Editor Errors" if team == "video_editing" else "Total Errors", "Errors per Article"]
    if team == "video_editing":
        headers.insert(-1, "Total QC Items")
    _style_header(ws, headers)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 27

    eligible_ids = _team_trend_eligible_article_ids(projects, bundles, error_rows)
    by_article = _article_error_map(error_rows, team, performance_only=False)
    grouped: dict[str, dict[str, Any]] = defaultdict(
        lambda: {"articles": set(), **{category: 0 for category in categories}}
    )

    for project in projects:
        article_id = str(project.get("article_id") or "")
        if article_id not in eligible_ids:
            continue
        month = _month_key(project.get(cfg["date_key"]))
        grouped[month]["articles"].add(article_id)
        for row in by_article.get(article_id, []):
            category = row.get("ai_category")
            if category in categories:
                grouped[month][category] += int(row.get("ai_error_count") or 0)

    for month in sorted(grouped):
        data = grouped[month]
        counts = [data[category] for category in categories]
        total_qc = sum(counts)
        total_perf = (
            sum(data[category] for category in VIDEO_CATEGORIES)
            if team == "video_editing"
            else total_qc
        )
        articles = len(data["articles"])
        row = [month, articles, *counts, total_perf]
        if team == "video_editing":
            row.append(total_qc)
        row.append(total_perf / articles if articles else 0)
        ws.append(row)

    if ws.max_row > 1:
        total_col = 3 + len(categories)
        ws.cell(
            ws.max_row + 1,
            total_col,
            sum(ws.cell(row, total_col).value or 0 for row in range(2, ws.max_row + 1)),
        )
        for row in ws.iter_rows(min_row=2, max_col=len(headers)):
            for cell in row[1:-1]:
                cell.alignment = Alignment(horizontal="center")
            row[-1].number_format = "0.00"

        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.height = 8
        chart.width = 10
        data = Reference(ws, min_col=total_col, min_row=1, max_row=ws.max_row - 1)
        categories_ref = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row - 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories_ref)
        chart.legend = None
        chart.y_axis.title = "AI-Classified Errors"
        chart.x_axis.title = cfg["trend_month_header"]
        ws.add_chart(chart, "I1")

    return ws


def _performance_project_counts(projects, bundles, error_rows, team):
    eligible_ids = _performance_eligible_article_ids(projects, bundles, error_rows, team)
    counts = Counter()
    for project in projects:
        article_id = str(project.get("article_id") or "")
        if article_id not in eligible_ids:
            continue
        person = _article_contributor(project, team)
        if person and person != "Needs Review":
            counts[person] += 1
    return counts


def _performance_errors(error_rows, team, eligible_ids: set[str] | None = None):
    rows = [
        row
        for row in error_rows
        if row.get("team") == team
        and row.get("performance_eligible", False)
        and row.get("ai_assignee")
    ]
    if eligible_ids is not None:
        rows = [
            row for row in rows
            if str(row.get("article_id") or "") in eligible_ids
        ]
    return rows


def _summary_sheet(wb, projects, bundles, error_rows, team):
    cfg = _team_config(team)
    categories = _performance_category_config(team)
    ws = wb.create_sheet(cfg["summary"])
    headers = [cfg["contributor"], "Articles", *categories.keys(), "Total Errors", "Errors per Article"]
    _style_header(ws, headers)

    eligible_ids = _performance_eligible_article_ids(projects, bundles, error_rows, team)
    project_counts = _performance_project_counts(projects, bundles, error_rows, team)
    category_counts = defaultdict(Counter)

    for row in _performance_errors(error_rows, team, eligible_ids):
        category_counts[str(row.get("ai_assignee"))][str(row.get("ai_category"))] += int(
            row.get("ai_error_count") or 0
        )

    rows = []
    for person, articles in project_counts.items():
        counts = [category_counts[person][category] for category in categories]
        total = sum(counts)
        rows.append([person, articles, *counts, total, total / articles if articles else 0])

    rows.sort(key=lambda row: (-row[-2], row[0]))
    for row in rows:
        ws.append(row)
        row_number = ws.max_row
        if int(row[1]) < 5:
            ws.cell(row_number, 2).fill = CAUTION_FILL
            ws.cell(row_number, 2).comment = Comment(
                "Small sample: fewer than 5 performance-eligible articles. Interpret Errors per Article cautiously.",
                "QC Automation",
            )

    for column in ws.iter_cols(
        min_col=len(headers), max_col=len(headers), min_row=2, max_row=ws.max_row
    ):
        for cell in column:
            cell.number_format = "0.00"

    return ws


def _contributor_monthly_sheet(wb, projects, bundles, error_rows, team):
    cfg = _team_config(team)
    categories = _performance_category_config(team)
    ws = wb.create_sheet(cfg["contributor_trend"])
    headers = [cfg["contributor"], cfg["trend_month_header"], "Performance-Eligible Articles", *categories.keys(), "Total Errors", "Errors per Article"]
    _style_header(ws, headers)

    project_map = {str(project.get("article_id") or ""): project for project in projects}
    eligible_ids = _performance_eligible_article_ids(projects, bundles, error_rows, team)
    project_counts = defaultdict(set)

    for project in projects:
        article_id = str(project.get("article_id") or "")
        if article_id not in eligible_ids:
            continue
        person = _article_contributor(project, team)
        if person and person != "Needs Review":
            key = (person, _month_key(project.get(cfg["date_key"])))
            project_counts[key].add(article_id)

    category_counts = defaultdict(Counter)
    for row in _performance_errors(error_rows, team, eligible_ids):
        project = project_map.get(str(row.get("article_id") or ""), {})
        key = (
            str(row.get("ai_assignee") or ""),
            _month_key(project.get(cfg["date_key"])),
        )
        category_counts[key][str(row.get("ai_category"))] += int(
            row.get("ai_error_count") or 0
        )

    ranges = {}
    contributors = sorted({key[0] for key in project_counts})
    cursor = 2

    for index, contributor in enumerate(contributors):
        start = cursor
        keys = sorted(
            [key for key in project_counts if key[0] == contributor],
            key=lambda key: key[1],
        )
        for key in keys:
            counts = [category_counts[key][category] for category in categories]
            total = sum(counts)
            articles = len(project_counts[key])
            ws.append(
                [
                    contributor,
                    key[1],
                    articles,
                    *counts,
                    total,
                    total / articles if articles else 0,
                ]
            )
            cursor += 1
        ranges[contributor] = (start, cursor - 1)
        if index != len(contributors) - 1:
            ws.append([None] * len(headers))
            cursor += 1

    for column in ws.iter_cols(
        min_col=len(headers), max_col=len(headers), min_row=2, max_row=ws.max_row
    ):
        for cell in column:
            cell.number_format = "0.00"

    candidates = []
    total_col = 4 + len(categories)
    metric_col = total_col + 1

    for person, (start, end) in ranges.items():
        if end - start + 1 >= 2:
            latest = str(ws.cell(end, 2).value or "")
            total = sum(int(ws.cell(row, total_col).value or 0) for row in range(start, end + 1))
            candidates.append((latest, total, person, start, end))

    candidates.sort(key=lambda item: (item[0], item[1]), reverse=True)
    positions = ["I6", "I21", "M6", "M21"]
    title_alias = {
        "Pallavi Sharma": "Pallavi",
        "Poornima G": "Poornima",
        "Kathyayani Sridharan": "KS",
        "Sri Lakshmi Priya": "SLP",
    }

    for position, (_, _, person, start, end) in zip(positions, candidates[:4]):
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.height = 7.5
        chart.width = 8
        data = Reference(ws, min_col=metric_col, min_row=start, max_row=end)
        categories_ref = Reference(ws, min_col=2, min_row=start, max_row=end)
        chart.add_data(data, titles_from_data=False)
        chart.set_categories(categories_ref)
        chart.title = title_alias.get(person, person)
        chart.legend = None
        chart.y_axis.title = "Errors per Article"
        ws.add_chart(chart, position)

    return ws


def _analysis_sheet(wb, projects, bundles, error_rows, team):
    categories = _category_config(team)
    ws = wb.create_sheet("Analysis Guide")
    _style_header(ws, ["Analysis Item", "Result / Interpretation"])
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 110

    totals = Counter()
    for row in error_rows:
        if _row_in_report(row, team):
            totals[str(row.get("ai_category"))] += int(row.get("ai_error_count") or 0)

    contributor_eligible_ids = _performance_eligible_article_ids(projects, bundles, error_rows, team)
    team_eligible_ids = _team_trend_eligible_article_ids(projects, bundles, error_rows)
    editor_or_writer_total = sum(totals[category] for category in _performance_category_config(team))
    vo_redos = totals["VO Redos"] if team == "video_editing" else 0

    rows = [("AI-classified performance-error total", editor_or_writer_total)]
    for category in categories:
        rows.append((category, totals[category]))

    rows.extend([
        ("Operational VO Redos excluded from editor performance", vo_redos),
        ("Team-trend verified articles", len(team_eligible_ids)),
        ("Contributor-performance eligible articles", len(contributor_eligible_ids)),
        ("Articles excluded from contributor denominator", len(projects) - len(contributor_eligible_ids)),
        (
            "Frame.io version rule",
            "Frame.io v1 is always excluded. Every version after v1 is checked. The latest version is included when it contains comments; it is excluded as the final/published version only when it has no comments. Filename text such as R, P, P2 or P3 is never used to decide inclusion.",
        ),
        (
            "Performance eligibility",
            "Team-level trends require verified extraction across every Frame.io version after v1 required by the current rule. Individual contributor performance additionally requires an unambiguous contributor.",
        ),
        (
            "VO Redos",
            "VO Redos are tracked visibly as operational QC items and are not assigned to or counted against a video editor. If the same source comment also contains independent video-editing instructions, those remain separate video-editing errors.",
        ),
        (
            "Reply counting",
            "A reply that only continues the same parent correction is not counted as a second error. The reply remains visible in Comment Coverage and is linked to the parent thread.",
        ),
        (
            "Comment preservation",
            "Every displayed source comment is copied verbatim from Frame.io. Comment Coverage shows every included QC comment, including comments classified as no error or consolidated into a parent thread.",
        ),
        (
            "Manual review workflow",
            "Use Review Overrides to approve, change, or remove AI-classified errors. Ambiguous editorial/scope exclusions are surfaced in Needs Review rather than silently finalized as non-performance.",
        ),
    ])

    for row in rows:
        ws.append(row)
    for row_number in range(2, ws.max_row + 1):
        ws.row_dimensions[row_number].height = 42
        ws.cell(row_number, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row_number, 2).alignment = Alignment(wrap_text=True, vertical="top")
    return ws

def _extraction_audit_sheet(wb, bundles):
    ws = wb.create_sheet("Extraction Audit")
    headers = [
        "Article Number",
        "Frame.io Review Link",
        "Resolution Status",
        "Resolution Method",
        "Asset Name",
        "Total Versions",
        "First Version (Always Excluded)",
        "Versions After v1 to Check",
        "Checked Versions After v1",
        "Included QC Versions",
        "Failed Versions After v1",
        "Checked Version Names",
        "Included Version Names",
        "Failed Version Names",
        "Latest Version",
        "Latest Included?",
        "Final/No-Comment Latest (Excluded)",
        "Extracted Comments/Replies",
        "Zero-Error Verified?",
        "Extraction Status",
        "Audit Note",
    ]
    _style_header(ws, headers)

    for bundle in sorted(bundles, key=lambda bundle: str((bundle.get("project") or {}).get("article_id") or "")):
        project = bundle.get("project") or {}
        resolution = bundle.get("resolution")
        audit = bundle.get("extraction_audit") or {}
        ws.append([
            project.get("item_name") or project.get("article_id"),
            project.get("frameio_review_link") or "",
            getattr(resolution, "status", "") if resolution else "",
            getattr(resolution, "method", "") if resolution else "",
            getattr(resolution, "asset_name", "") if resolution else "",
            audit.get("total_versions", 0),
            audit.get("first_version", ""),
            audit.get("eligible_versions", 0),
            audit.get("checked_versions", 0),
            audit.get("included_versions", 0),
            audit.get("failed_versions", 0),
            "; ".join(audit.get("checked_version_names") or []),
            "; ".join(audit.get("included_version_names") or []),
            "; ".join(audit.get("failed_version_names") or []),
            audit.get("latest_version", ""),
            "Yes" if audit.get("latest_version_included") else "No",
            audit.get("final_version", ""),
            audit.get("comment_count", 0),
            "Yes" if audit.get("zero_error_verified") else "No",
            audit.get("status", "needs_review"),
            audit.get("note", ""),
        ])
        _make_link(ws.cell(ws.max_row, 2), str(project.get("frameio_review_link") or ""))
        if audit.get("status") != "verified":
            for column in range(1, len(headers) + 1):
                ws.cell(ws.max_row, column).fill = CAUTION_FILL

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    return ws

def _needs_review_sheet(wb, bundles, error_rows, team):
    ws = wb.create_sheet("Needs Review")
    headers = [
        "Article Number", "Type", "Reason", "Frame.io Version", "Timecode",
        "Frame.io Review Link", "Comment", "AI Category", "AI Assignee",
        "Confidence", "Performance Counted?"
    ]
    _style_header(ws, headers)

    for bundle in bundles:
        project = bundle.get("project") or {}
        resolution = bundle.get("resolution")
        audit = bundle.get("extraction_audit") or {}
        if (
            resolution is None
            or getattr(resolution, "status", "") != "resolved"
            or audit.get("status") != "verified"
        ):
            reason = audit.get("note") or getattr(resolution, "note", "") or "Extraction requires review."
            ws.append([
                project.get("item_name") or project.get("article_id"),
                "Extraction Verification", reason, "", "",
                project.get("frameio_review_link") or "", "", "", "", "", "No"
            ])
            _make_link(ws.cell(ws.max_row, 6), str(project.get("frameio_review_link") or ""))

        if team == "scripting" and not str(project.get("scriptwriter") or "").strip():
            ws.append([
                project.get("item_name") or project.get("article_id"),
                "Monday Attribution",
                "Scriptwriter is blank in Monday.com. Excluded from individual writer performance until corrected.",
                "", "", project.get("frameio_review_link") or "", "", "", "", "", "No",
            ])
            _make_link(ws.cell(ws.max_row, 6), str(project.get("frameio_review_link") or ""))

        if team == "video_editing":
            editors = _distinct_editors(project)
            if not editors:
                ws.append([
                    project.get("item_name") or project.get("article_id"),
                    "Monday Attribution",
                    "No video editor is assigned in Monday.com. Excluded from individual editor performance until corrected.",
                    "", "", project.get("frameio_review_link") or "", "", "", "", "", "No",
                ])
                _make_link(ws.cell(ws.max_row, 6), str(project.get("frameio_review_link") or ""))
            elif len(editors) > 1:
                ws.append([
                    project.get("item_name") or project.get("article_id"),
                    "Editor Handover",
                    "Multiple video editors are assigned in Monday.com: " + ", ".join(editors) + ". Individual error attribution requires review.",
                    "", "", project.get("frameio_review_link") or "", "", "", "", "", "No",
                ])
                _make_link(ws.cell(ws.max_row, 6), str(project.get("frameio_review_link") or ""))

    for row in error_rows:
        row_team = row.get("team")
        operational_review = (
            team == "video_editing"
            and row_team == "operational"
            and bool(row.get("needs_review"))
        )
        if row_team not in {team, "review"} and not operational_review:
            continue
        if row.get("needs_review") or not row.get("performance_eligible", True):
            if operational_review:
                review_type = "Operational Scope Review"
            else:
                review_type = (
                    "Classification Review"
                    if row.get("needs_review") and row.get("performance_eligible", False)
                    else "Classification / Attribution"
                )
            ws.append([
                row.get("item_name") or row.get("article_id"),
                review_type,
                row.get("classification_reason") or "Needs manual verification",
                row.get("version_label") or (f"v{row.get('version_number')}" if row.get("version_number") else ""),
                row.get("timecode", ""),
                row.get("frameio_review_link", ""),
                row.get("comment_text") or "",
                row.get("ai_category") or "",
                row.get("ai_assignee") or "",
                row.get("confidence", ""),
                "Yes" if row.get("performance_eligible", False) else "No",
            ])
            _make_link(ws.cell(ws.max_row, 6), str(row.get("frameio_review_link") or ""))

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    return ws

def _error_key(row: dict[str, Any]) -> str:
    if row.get("source_error_key"):
        return str(row["source_error_key"])
    return f"{row.get('comment_id', '')}:{row.get('issue_index', 0)}:{row.get('team', '')}"


def _validation_lists_sheet(wb):
    ws = wb.create_sheet("_Validation Lists")
    ws.append(["Review Actions", "Teams", "Categories"])
    actions = ["Approve", "Change", "Remove"]
    teams = ["scripting", "video_editing", "operational", "none"]
    categories = list(SCRIPTING_CATEGORIES) + list(VIDEO_CATEGORIES) + list(OPERATIONAL_CATEGORIES) + ["None"]
    max_rows = max(len(actions), len(teams), len(categories))
    for index in range(max_rows):
        ws.append([
            actions[index] if index < len(actions) else "",
            teams[index] if index < len(teams) else "",
            categories[index] if index < len(categories) else "",
        ])
    ws.sheet_state = "hidden"
    return {
        "actions": "'_Validation Lists'!$A$2:$A$" + str(len(actions) + 1),
        "teams": "'_Validation Lists'!$B$2:$B$" + str(len(teams) + 1),
        "categories": "'_Validation Lists'!$C$2:$C$" + str(len(categories) + 1),
    }

def _review_overrides_sheet(wb, error_rows, team):
    validation_ranges = _validation_lists_sheet(wb)
    ws = wb.create_sheet("Review Overrides")
    headers = [
        "Error Key",
        "Article Number",
        "Version",
        "Timecode",
        "Frame.io Review Link",
        "Comment",
        "AI Team",
        "AI Category",
        "AI Assignee",
        "AI Count",
        "AI Needs Review",
        "AI Performance Eligible",
        "Review Action",
        "Manual Team",
        "Manual Category",
        "Manual Assignee",
        "Manual Count",
        "Reviewer Notes",
    ]
    _style_header(ws, headers)

    rows = [
        row
        for row in error_rows
        if (
            row.get("team") in {team, "review"}
            or (team == "video_editing" and row.get("team") == "operational")
        )
    ]
    rows.sort(
        key=lambda row: (
            0 if row.get("needs_review") or not row.get("performance_eligible", True) else 1,
            str(row.get("article_id") or ""),
            int(row.get("version_number") or 0),
            str(row.get("comment_id") or ""),
            int(row.get("issue_index") or 0),
        )
    )

    for row in rows:
        ws.append(
            [
                _error_key(row),
                row.get("item_name") or row.get("article_id"),
                row.get("version_number", ""),
                row.get("timecode", ""),
                row.get("frameio_review_link", ""),
                row.get("comment_text", ""),
                row.get("team", ""),
                row.get("ai_category", ""),
                row.get("ai_assignee", ""),
                row.get("ai_error_count", 0),
                "Yes" if row.get("needs_review") else "No",
                "Yes" if row.get("performance_eligible", False) else "No",
                row.get("review_action", ""),
                row.get("manual_team", ""),
                row.get("manual_category", ""),
                row.get("manual_assignee", ""),
                row.get("manual_count", "") if row.get("manual_count") is not None else "",
                row.get("reviewer_notes", ""),
            ]
        )
        row_number = ws.max_row
        _make_link(ws.cell(row_number, 5), str(row.get("frameio_review_link") or ""))
        for column in range(1, 13):
            ws.cell(row_number, column).font = IMPORTED_FONT
        for column in range(13, 19):
            ws.cell(row_number, column).font = INPUT_FONT

    action_validation = DataValidation(type="list", formula1=validation_ranges["actions"], allow_blank=True)
    team_validation = DataValidation(type="list", formula1=validation_ranges["teams"], allow_blank=True)
    category_validation = DataValidation(
        type="list",
        formula1=validation_ranges["categories"],
        allow_blank=True,
    )
    ws.add_data_validation(action_validation)
    ws.add_data_validation(team_validation)
    ws.add_data_validation(category_validation)
    if ws.max_row >= 2:
        action_validation.add(f"M2:M{ws.max_row}")
        team_validation.add(f"N2:N{ws.max_row}")
        category_validation.add(f"O2:O{ws.max_row}")

    widths = {
        "A": 52, "B": 24, "C": 10, "D": 16, "E": 30, "F": 80,
        "G": 18, "H": 30, "I": 25, "J": 10, "K": 16, "L": 22,
        "M": 18, "N": 20, "O": 32, "P": 25, "Q": 14, "R": 60,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"
    return ws


def _error_detail_sheet(wb, error_rows, team):
    ws = wb.create_sheet("Error Detail")
    headers = [
        "Error Key", "Article Number", "Team", "Category", "Error Count", "Source Comment",
        "AI Assignee", "Performance Eligible", "Needs Review", "Confidence",
        "Frame.io Version", "Version File", "Timecode", "Frame.io Review Link",
        "Comment ID", "Pattern Label",
    ]
    _style_header(ws, headers)

    for row in error_rows:
        if not (_row_in_report(row, team) or row.get("team") == "review"):
            continue
        ws.append([
            _error_key(row),
            row.get("item_name") or row.get("article_id"),
            row.get("team"),
            row.get("ai_category"),
            row.get("ai_error_count"),
            row.get("comment_text", ""),
            row.get("ai_assignee"),
            "Yes" if row.get("performance_eligible", False) else "No",
            "Yes" if row.get("needs_review") else "No",
            row.get("confidence", ""),
            row.get("version_label") or (f"v{row.get('version_number')}" if row.get("version_number") else ""),
            row.get("version_name", ""),
            row.get("timecode", ""),
            row.get("frameio_review_link", ""),
            row.get("comment_id", ""),
            row.get("pattern_label", ""),
        ])
        _make_link(ws.cell(ws.max_row, 14), str(row.get("frameio_review_link") or ""))

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.sheet_state = "hidden"
    return ws

def _raw_comments_sheet(wb, bundles):
    ws = wb.create_sheet("Raw Frame.io Comments")
    headers = [
        "Article Number", "Frame.io Version", "Version File", "Comment ID",
        "Parent Comment ID", "Reply?", "Commenter", "Created At", "Timecode",
        "Comment", "Repeat Match?", "Prior Match Version", "Prior Match Score", "Prior Match Text",
    ]
    _style_header(ws, headers)

    for bundle in bundles:
        project = bundle.get("project") or {}
        for row in bundle.get("comments") or []:
            ws.append([
                project.get("item_name") or row.get("article_id"),
                row.get("version_label") or (f"v{row.get('version_number')}" if row.get("version_number") else ""),
                row.get("version_name"),
                row.get("comment_id"),
                row.get("parent_comment_id"),
                "Yes" if row.get("is_reply") else "No",
                row.get("commenter"),
                row.get("created_at"),
                row.get("timestamp"),
                row.get("text"),
                "Yes" if row.get("repeated_from_prior_version") else "No",
                row.get("prior_match_version", ""),
                row.get("prior_match_score", ""),
                row.get("prior_match_text", ""),
            ])
    ws.sheet_state = "hidden"
    return ws

def _canonical_pattern(category: str, pattern: str) -> str:
    label = (pattern or category or "").strip()
    low = label.casefold()

    if category == "Previous Comments Unaddressed":
        return "Prior comment unaddressed"

    if category == "Audio/Visual":
        if any(token in low for token in ("audio-video", "footage-vo", "vo mismatch", "audio visual mismatch", "sync")):
            return "Audio-video mismatch"
        if "pacing" in low or "pace" in low:
            return "Footage pacing"
        if "blurr" in low:
            return "Blurry footage"
        if "sound" in low and "glitch" in low:
            return "Sound glitch"
        if "video" in low and "glitch" in low:
            return "Video glitch"

    if category == "On-Screen Text":
        if "spell" in low or "typo" in low:
            return "Spelling error"
        if any(token in low for token in ("format", "font", "italic", "capital", "style")):
            return "Format error"
        if any(token in low for token in ("interpret", "incorrect text", "wrong text")):
            return "Interpretation error"

    if category == "Result Section":
        if "highlight" in low:
            return "Highlighting issue"
        if "label" in low:
            return "Labelling issue"
        if "contrast" in low or "colour" in low or "color" in low:
            return "Contrast colour issue"

    return label


def _pattern_sheet(wb, error_rows):
    ws = wb.create_sheet("Common Error Patterns")
    _style_header(ws, ["Category", "Pattern", "Total Count", "Articles", "Editors", "Recurring?"])

    grouped = defaultdict(lambda: {"count": 0, "articles": set(), "editors": set()})
    for row in error_rows:
        if row.get("team") != "video_editing":
            continue
        category = str(row.get("ai_category") or "")
        pattern = _canonical_pattern(category, str(row.get("pattern_label") or category))
        key = (category, pattern)
        grouped[key]["count"] += int(row.get("ai_error_count") or 0)
        if row.get("article_id"):
            grouped[key]["articles"].add(row.get("article_id"))
        if row.get("ai_assignee"):
            grouped[key]["editors"].add(row.get("ai_assignee"))

    for (category, pattern), info in sorted(
        grouped.items(), key=lambda item: (-item[1]["count"], item[0])
    ):
        recurring = info["count"] >= 2 and len(info["articles"]) >= 2
        ws.append(
            [
                category,
                pattern,
                info["count"],
                len(info["articles"]),
                len(info["editors"]),
                "Yes" if recurring else "No",
            ]
        )
    return ws


def _scope_reconciliation_sheet(wb, projects, bundles, error_rows):
    ws = wb.create_sheet("Scope Reconciliation")
    _style_header(ws, ["Metric", "Count", "Interpretation"])

    verified_articles = {
        str((bundle.get("project") or {}).get("article_id") or "")
        for bundle in bundles
        if _extraction_verified(bundle)
    }
    contributor_eligible = _performance_eligible_article_ids(
        projects, bundles, error_rows, "video_editing"
    )
    editor_errors = sum(
        int(row.get("ai_error_count") or 0)
        for row in error_rows
        if row.get("team") == "video_editing"
    )
    vo_redos = sum(
        int(row.get("ai_error_count") or 0)
        for row in error_rows
        if row.get("team") == "operational" and row.get("ai_category") == "VO Redos"
    )
    pattern_total = sum(
        int(row.get("ai_error_count") or 0)
        for row in error_rows
        if row.get("team") == "video_editing"
    )
    review_items = (
        wb["Needs Review"].max_row - 1
        if "Needs Review" in wb.sheetnames
        else 0
    )

    rows = [
        ("Monthly Monday projects", len(projects), "Complete selected release-month population."),
        ("Verified extraction articles", len(verified_articles), "Articles for which every Frame.io version after v1 required by the current rule was checked successfully."),
        ("Individual-editor eligible articles", len(contributor_eligible), "Verified articles with an unambiguous video editor and no blocking attribution issue."),
        ("Video-editor errors", editor_errors, "Errors attributed to video editing; this is the total used by editor performance summaries."),
        ("VO Redos (non-editor)", vo_redos, "Operational VO redo requests shown separately and excluded from editor performance."),
        ("Total QC items", editor_errors + vo_redos, "Video-editor errors plus operational VO Redos."),
        ("Common Error Patterns total", pattern_total, "All video-editor error patterns are listed, including one-off patterns, so this reconciles to Video-editor errors."),
        ("Review items", review_items, "Exact row count from the Video Needs Review sheet."),
    ]
    for row in rows:
        ws.append(row)

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 100
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    return ws

def _stage_breakdown_sheet(wb, error_rows):
    ws = wb.create_sheet("Frame.io Version Breakdown")
    categories = {**VIDEO_CATEGORIES, "VO Redos": OPERATIONAL_CATEGORIES["VO Redos"]}
    _style_header(ws, ["Frame.io Version", "Articles", *categories.keys(), "Total Editor Errors", "Total QC Items"])
    grouped = defaultdict(lambda: {"articles": set(), **{category: 0 for category in categories}})
    for row in error_rows:
        if not _row_in_report(row, "video_editing"):
            continue
        version_label = str(row.get("version_label") or (f"v{row.get('version_number')}" if row.get("version_number") else "Unknown"))
        grouped[version_label]["articles"].add(str(row.get("article_id") or ""))
        category = str(row.get("ai_category") or "")
        if category in categories:
            grouped[version_label][category] += int(row.get("ai_error_count") or 0)
    for version_label, data in sorted(grouped.items()):
        counts = [data[category] for category in categories]
        editor_total = sum(data[category] for category in VIDEO_CATEGORIES)
        ws.append([version_label, len(data["articles"]), *counts, editor_total, sum(counts)])
    return ws

def _comment_coverage_sheet(wb, bundles, classifications):
    ws = wb.create_sheet("Comment Coverage")
    _style_header(ws, ["Article Number", "Frame.io Version", "Version File", "Commenter", "Timecode", "Source Comment", "Classification Outcome"])
    by_comment = {str(row.get("comment_id") or ""): row for row in classifications}
    for bundle in bundles:
        project = bundle.get("project") or {}
        for comment in bundle.get("comments") or []:
            result = by_comment.get(str(comment.get("comment_id") or ""), {})
            issues = result.get("issues") or []
            dedup_parent = str(result.get("_deduplicated_to_parent") or "")
            if result.get("_classification_error"):
                outcome = "Needs Review - classification failure"
            elif dedup_parent and not issues:
                outcome = f"Thread continuation - counted with parent comment {dedup_parent}"
            elif not issues:
                outcome = "No counted QC error"
            else:
                outcome = "; ".join(
                    f"{issue.get('team')}: {issue.get('category')} x{int(issue.get('error_count') or 0)}"
                    for issue in issues
                )
                if dedup_parent:
                    outcome += f"; duplicate issue consolidated with parent comment {dedup_parent}"
            ws.append([
                project.get("item_name") or comment.get("article_id"),
                comment.get("version_label") or (f"v{comment.get('version_number')}" if comment.get("version_number") else ""),
                comment.get("version_name", ""),
                comment.get("commenter", ""),
                comment.get("timestamp", ""),
                comment.get("text", ""),
                outcome,
            ])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    return ws

def _non_performance_items_sheet(wb, error_rows):
    ws = wb.create_sheet("Non-Performance Items")
    _style_header(ws, ["Article Number", "Frame.io Version", "Category", "Source Comment", "Reason", "Timecode", "Frame.io Review Link"])
    for row in error_rows:
        if row.get("team") != "operational":
            continue
        ws.append([
            row.get("item_name") or row.get("article_id"),
            row.get("version_label") or (f"v{row.get('version_number')}" if row.get("version_number") else ""),
            row.get("ai_category", ""),
            row.get("comment_text", ""),
            row.get("classification_reason", ""),
            row.get("timecode", ""),
            row.get("frameio_review_link", ""),
        ])
        _make_link(ws.cell(ws.max_row, 7), str(row.get("frameio_review_link") or ""))
    return ws

def _excluded_source_sheet(wb, bundles):
    ws = wb.create_sheet("Excluded Version Policy")
    _style_header(ws, ["Article Number", "Excluded Frame.io Version", "Version File", "Reason"])
    for bundle in bundles:
        project = bundle.get("project") or {}
        versions = bundle.get("versions") or []
        for version in versions:
            if version.get("included_for_qc"):
                continue
            ws.append([
                project.get("item_name") or project.get("article_id"),
                version.get("version_label") or (f"v{version.get('version_number')}" if version.get("version_number") else ""),
                version.get("name") or "",
                version.get("qc_check_note") or "Excluded by the current Frame.io version rule.",
            ])
    return ws

def _build_info_sheet(wb, team: str, year: int, month: int, review_applied: bool) -> None:
    ws = wb.create_sheet("_Build Info")
    ws.append(["Build Version", BUILD_VERSION])
    ws.append(["Build Label", BUILD_LABEL])
    ws.append(["Report Team", team])
    ws.append(["Selected Release Month", f"{year:04d}-{month:02d}"])
    ws.append(["Review Overrides Applied", "Yes" if review_applied else "No"])
    ws.append(["Reference", "Swati June QC format + approved Video QC error sheet"])
    ws.sheet_state = "hidden"


def build_report(
    *,
    team: str,
    year: int,
    month: int,
    projects: list[dict[str, Any]],
    bundles: list[dict[str, Any]],
    error_rows: list[dict[str, Any]],
    classifications: list[dict[str, Any]] | None = None,
    review_applied: bool = False,
) -> bytes:
    if team not in {"scripting", "video_editing"}:
        raise ValueError("team must be scripting or video_editing")

    wb = Workbook()
    wb.remove(wb.active)

    _main_sheet(wb, projects, bundles, error_rows, team)
    _description_sheet(wb, team)
    _monthly_trend_sheet(wb, projects, bundles, error_rows, team)
    _summary_sheet(wb, projects, bundles, error_rows, team)
    _contributor_monthly_sheet(wb, projects, bundles, error_rows, team)

    if team == "video_editing":
        _pattern_sheet(wb, error_rows)
        _stage_breakdown_sheet(wb, error_rows)

    _comment_coverage_sheet(wb, bundles, classifications or [])
    _non_performance_items_sheet(wb, error_rows)
    _excluded_source_sheet(wb, bundles)
    _analysis_sheet(wb, projects, bundles, error_rows, team)
    _extraction_audit_sheet(wb, bundles)
    _needs_review_sheet(wb, bundles, error_rows, team)
    if team == "video_editing":
        _scope_reconciliation_sheet(wb, projects, bundles, error_rows)
    _review_overrides_sheet(wb, error_rows, team)
    _error_detail_sheet(wb, error_rows, team)
    _raw_comments_sheet(wb, bundles)
    _build_info_sheet(wb, team, year, month, review_applied)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
